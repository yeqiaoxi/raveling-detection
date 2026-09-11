from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT=Path(r"Y:\jiaotong\output\docx\CFDSLite_UNet_PaveSeg多病害训练指南.docx"); OUT.parent.mkdir(parents=True,exist_ok=True)
d=Document(); s=d.sections[0]; s.top_margin=Cm(2); s.bottom_margin=Cm(1.9); s.left_margin=Cm(2.2); s.right_margin=Cm(2.1)
for n,z in [("Normal",10.5),("Title",23),("Subtitle",11),("Heading 1",16),("Heading 2",12.5)]:
 st=d.styles[n]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); st.font.size=Pt(z); st.font.color.rgb=RGBColor(0,0,0)
d.styles["Normal"].paragraph_format.line_spacing=1.22; d.styles["Normal"].paragraph_format.space_after=Pt(6)
for n in ("Heading 1","Heading 2"): d.styles[n].paragraph_format.keep_with_next=True; d.styles[n].paragraph_format.space_before=Pt(13); d.styles[n].paragraph_format.space_after=Pt(6)

def shade(c,color):
 p=c._tc.get_or_add_tcPr(); x=OxmlElement("w:shd"); x.set(qn("w:fill"),color); p.append(x)
def tab(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 pr=t._tbl.tblPr; b=OxmlElement("w:tblBorders")
 for e in ("top","left","bottom","right","insideH","insideV"):
  x=OxmlElement(f"w:{e}"); x.set(qn("w:val"),"single"); x.set(qn("w:sz"),"4"); x.set(qn("w:color"),"D9D9D9"); b.append(x)
 pr.append(b); hp=t.rows[0]._tr.get_or_add_trPr(); rr=OxmlElement("w:tblHeader"); rr.set(qn("w:val"),"true"); hp.append(rr)
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=h; c.width=Cm(widths[i]); shade(c,"1F4E78")
  for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
 for ri,row in enumerate(rows):
  cs=t.add_row().cells
  for i,v in enumerate(row):
   cs[i].text=str(v); cs[i].width=Cm(widths[i]); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if ri%2: shade(cs[i],"F2F6FA")
   for p in cs[i].paragraphs:
    p.paragraph_format.space_after=Pt(1)
    for r in p.runs:r.font.size=Pt(8.7)
 d.add_paragraph()
def code(x):
 p=d.add_paragraph(); p.paragraph_format.left_indent=Cm(.4); p.paragraph_format.space_after=Pt(7); sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),"F3F5F7"); p._p.get_or_add_pPr().append(sh)
 for line in x.strip().splitlines():
  r=p.add_run(line+"\n"); r.font.name="Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); r.font.size=Pt(8.4)
def bullet(x): d.add_paragraph(x,style="List Bullet")

p=d.add_paragraph(style="Title"); p.add_run("CFDSLite UNet PaveSeg多病害训练指南")
d.add_paragraph("基于轻量化CFDSLite UNet的路面多类型病害语义分割研究",style="Subtitle")
d.add_paragraph("适用环境：摩禾云 V100 32 GB    数据：PaveSeg公开预览集    版本：1.0",style="Subtitle")
d.add_paragraph("本指南全部在GPU服务器终端执行，从GitHub更新代码、下载PaveSeg、核验八类别掩膜，到训练、预测和结果备份。PaveSeg公开仓库当前提供100张预览图和像素级标注，因此可以完成流程验证和小样本实验，但不能把预览集结果等同于完整2400张数据集的正式基准结果。")

d.add_heading("一 研究任务与类别",1)
d.add_paragraph("模型执行八类别语义分割，每个像素只能属于一个类别。网络保留论文CFDSLite-UNet的深度可分离卷积和双线性上采样，将输出层从单通道改为八通道，并使用CrossEntropyLoss。")
tab(["类别ID","类别","PaveSeg灰度值"],[[0,"背景",0],[1,"裂缝",30],[2,"坑槽",60],[3,"封缝",90],[4,"修补",120],[5,"龟裂",150],[6,"井盖",180],[7,"伸缩缝",210]],[2.2,7.5,5.5])
d.add_paragraph("该数据集不包含松散病害。研究名称和结论应使用“多类型路面病害分割”，不能把裂缝、龟裂或坑槽改名为松散。")

d.add_heading("二 更新云端项目",1)
d.add_paragraph("先进入已经克隆的项目并获取最新多类别代码。若git pull提示本地文件冲突，不要强制覆盖，先保存终端输出。")
code("cd ~/raveling-detection\ngit pull origin main\nls\npython3 -m py_compile cfds_lite_unet.py paveseg_training.py")
d.add_paragraph("目录中必须出现 paveseg_training.py。激活原有虚拟环境并安装依赖：")
code("cd ~/raveling-detection\nsource .venv/bin/activate\npython -m pip install -r requirements.txt")

d.add_heading("三 下载PaveSeg",1)
d.add_paragraph("PaveSeg位于公开GitHub仓库，无需Kaggle账号，不使用任何占位下载链接。")
code("cd ~/raveling-detection\nmkdir -p data/raw\ncd data/raw\ngit clone --depth 1 https://github.com/FuturePave/PaveSeg-Dataset.git paveseg")
d.add_paragraph("如果目录已存在，则更新而不是重复克隆：")
code("cd ~/raveling-detection/data/raw/paveseg\ngit pull")
d.add_heading("3.1 核验数据",2)
code("cd ~/raveling-detection\nfind data/raw/paveseg/img_preview -type f | wc -l\nfind data/raw/paveseg/label_preview -type f | wc -l\nls data/raw/paveseg/img_preview | head\nls data/raw/paveseg/label_preview | head")
d.add_paragraph("公开预览版本应有100张图像和100张掩膜。配对格式为 000001.jpg 与 000001_mask.png。若数量不是100，先检查git clone是否完成。")
d.add_heading("3.2 检查掩膜灰度值",2)
code("python - <<'PY'\nimport cv2, glob, numpy as np\nvalues=set()\nfor f in glob.glob('data/raw/paveseg/label_preview/*.png'):\n    values.update(np.unique(cv2.imread(f,0)).tolist())\nprint(sorted(values))\nPY")
d.add_paragraph("正确输出应为 [0, 30, 60, 90, 120, 150, 180, 210] 的全部或子集。程序会自动转换为类别ID 0至7；发现其他灰度值时会停止，以防静默生成错误标签。")

d.add_heading("四 检查V100环境",1)
code("nvidia-smi\npython -c \"import torch; print(torch.__version__); print(torch.cuda.is_available()); print(torch.cuda.get_device_name(0) if torch.cuda.is_available() else 'NO GPU')\"")
d.add_paragraph("必须看到 True 和 Tesla V100。若为False，不要开始训练。V100 32 GB足以使用512×512、物理batch 4。")

d.add_heading("五 一轮试跑",1)
d.add_paragraph("先运行1个epoch，检查读取、标签转换、CUDA和权重保存。")
code("cd ~/raveling-detection\nsource .venv/bin/activate\nmkdir -p runs\npython -u paveseg_training.py train --data data/raw/paveseg --output runs/paveseg_smoke.pt --epochs 1 --batch-size 1 --workers 0 --device cuda")
bullet("应显示 device=cuda、train=80、val=20。")
bullet("参数量约3.55M，八类别输出层会比原单类别模型略多。")
bullet("runs/paveseg_smoke.pt 应存在且大小不为零。")

d.add_heading("六 正式训练",1)
d.add_heading("6.1 论文风格基线",2)
code("python -u paveseg_training.py train --data data/raw/paveseg --output runs/paveseg_best.pt --epochs 100 --batch-size 4 --accumulation-steps 1 --lr 1e-4 --val-ratio 0.2 --workers 2 --device cuda 2>&1 | tee runs/paveseg_train.log")
d.add_paragraph("该命令使用物理batch 4，便于与原论文训练设置保持接近。预览集只有100张，100轮可能过拟合，应同时观察验证mIoU和预测图。")
d.add_heading("6.2 有效Batch size 20",2)
d.add_paragraph("不建议直接设置物理batch 20。使用物理batch 4并累积5次梯度，可得到有效batch 20，显存占用仍接近batch 4。")
code("python -u paveseg_training.py train --data data/raw/paveseg --output runs/paveseg_bs20.pt --epochs 100 --batch-size 4 --accumulation-steps 5 --lr 1e-4 --val-ratio 0.2 --workers 2 --device cuda 2>&1 | tee runs/paveseg_bs20.log")
d.add_paragraph("有效batch 20是独立对照实验，不是论文原参数。100张预览数据每轮仅20个训练batch，累积5次后每轮只有约4次参数更新，收敛可能更慢。因此推荐先训练batch 4基线，再比较batch 20。")
d.add_heading("6.3 防止终端断线",2)
code("tmux new -s paveseg\n# 在tmux中运行训练命令\n# 离开：Ctrl+B，再按D\n# 返回：tmux attach -t paveseg")

d.add_heading("七 指标解释",1)
tab(["指标","用途","注意事项"],[["mIoU","八个类别IoU的平均值","某类在验证集不存在时忽略该类"],["Pixel Accuracy","所有预测正确像素比例","背景占比大时可能虚高"],["各类别IoU","分别评价裂缝、坑槽等","小样本类别波动很大"],["验证损失","观察收敛和过拟合","不能代替独立测试集"]],[3.2,5.5,6.5])
d.add_paragraph("当前代码按固定随机种子42进行80比20随机划分。正式论文实验应保存划分清单，并优先按道路或采集批次划分，避免相邻图像同时进入训练集和验证集。")

d.add_heading("八 预测与可视化",1)
code("python paveseg_training.py predict --input data/raw/paveseg/img_preview --weights runs/paveseg_best.pt --output runs/paveseg_predictions --device cuda")
d.add_paragraph("输出包括每张图片的彩色类别图、与原图叠加的可视化结果，以及 class_ratios.csv。CSV记录每类像素占整张图片的比例，可用于病害构成分析，但不能在没有尺度标定时换算为平方米。")
tab(["颜色","类别"],[["黑色","背景"],["红色","裂缝"],["橙色","坑槽"],["黄色","封缝"],["蓝色","修补"],["紫色","龟裂"],["青色","井盖"],["绿色","伸缩缝"]],[5,10.2])

d.add_heading("九 常见问题",1)
tab(["问题","处理"],[["No PaveSeg pairs found","确认--data指向同时包含img_preview和label_preview的PaveSeg根目录"],["Unknown grayscale label","掩膜被压缩或改色；使用原始PNG，不要保存为JPG"],["CUDA out of memory","依次把batch改为2或1；保持accumulation steps获得目标有效batch"],["训练只预测背景","类别极不平衡；检查标签值并增加类别权重或裁剪包含病害的区域"],["mIoU波动大","验证集只有20张且类别不均衡；多次分层实验或申请完整数据"],["git clone超时","设置git HTTP 1.1后重试，或使用平台网络加速；不要中途按Ctrl+C"]],[5.2,10])

d.add_heading("十 保存结果并停止计费",1)
code("cd ~/raveling-detection\ntar -czf paveseg_results.tar.gz runs/paveseg_best.pt runs/paveseg_train.log runs/paveseg_predictions\nls -lh paveseg_results.tar.gz")
d.add_paragraph("通过平台文件管理器下载压缩包。确认本地可打开后再停止GPU实例；训练进程结束不代表平台自动停止计费。不要把数据、权重或Kaggle密钥提交到Git仓库。")

d.add_heading("十一 实验报告建议",1)
bullet("报告每个类别的IoU，不只报告整体像素准确率。")
bullet("分别比较batch 4和有效batch 20，并记录训练时间与峰值显存。")
bullet("明确说明当前使用PaveSeg 100张公开预览集，结论属于小样本验证。")
bullet("将CFDSLite-UNet与基础U-Net进行参数量、mIoU和推理速度对比。")
bullet("不将多病害实验结果表述为松散检测结果。")

for sec in d.sections:
 f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.add_run("CFDSLite UNet PaveSeg训练指南    ")
 fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); f._p.append(fld)
 for r in f.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,100,100)
d.core_properties.title="CFDSLite UNet PaveSeg多病害训练指南"; d.core_properties.author=""
d.save(OUT); print(OUT)
